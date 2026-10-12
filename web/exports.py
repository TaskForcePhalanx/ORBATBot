"""Briefing packs as files — the PDF and Word renderers.

What a pack *says* is decided in `utils/campaign.document_blocks()`, under
test. This module only lays that list of blocks out on a page, twice: once
with reportlab (PDF, for attaching to Discord) and once with python-docx
(Word, which opens straight into Google Docs). Neither renderer decides
anything about content, so the two files always say the same thing.

Both libraries are imported inside the functions that use them, the way
`utils/tiles.py` treats py7zr: they are declared dependencies, but a deployment
that somehow lacks them keeps every other page working and gets a message here.

Branding is the site's own: `WEB_BRAND` and the logo in `web/static/`.
"""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

from web.config import LOGO_NAMES

_STATIC = Path(__file__).parent / 'static'
_INK = (0x1B, 0x20, 0x27)
_MUTED = (0x5D, 0x68, 0x74)
_ACCENT = (0xB8, 0x97, 0x3F)        # a muted gold rule — prints well in greyscale too


def logo_bytes() -> bytes | None:
    """The site logo as PNG/JPEG bytes, or None. An SVG logo is skipped: neither
    library draws SVG, and a cover without a logo beats a failed export."""
    for name in LOGO_NAMES:
        path = _STATIC / name
        if path.exists() and not name.endswith('.svg'):
            return path.read_bytes()
    return None


def _png_or_jpeg(data: bytes) -> bytes:
    """Word takes PNG, JPEG and GIF; WebP (and anything odd) is converted."""
    if data[:8] == b'\x89PNG\r\n\x1a\n' or data[:3] == b'\xff\xd8\xff':
        return data
    from PIL import Image
    out = io.BytesIO()
    with Image.open(io.BytesIO(data)) as im:
        im.convert('RGBA' if im.mode in ('RGBA', 'LA', 'P') else 'RGB').save(out, 'PNG')
    return out.getvalue()


def _size(data: bytes) -> tuple[int, int]:
    from PIL import Image
    with Image.open(io.BytesIO(data)) as im:
        return im.size


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

_FONTS_READY = False
# Characters Bitstream Vera (the TrueType font that ships with reportlab) has no
# glyph for, spelled with ones it has — so an arrow in an objective prints as
# an arrow-ish "->" rather than a black box.
_FALLBACKS = {'→': '->', '←': '<-', '↑': '^', '↓': 'v', '⇒': '=>', '✓': 'v', '✔': 'v',
              '✗': 'x', '✘': 'x', '•': '·', '−': '-', '≈': '~', '≤': '<=', '≥': '>=',
              ' ': ' ', ' ': ' ', ' ': ' '}


def _register_fonts():
    global _FONTS_READY, _CMAP
    if _FONTS_READY:
        return
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.fonts import addMapping
    base = Path(reportlab.__file__).parent / 'fonts'
    for name, file in (('Vera', 'Vera.ttf'), ('Vera-Bold', 'VeraBd.ttf'),
                       ('Vera-Italic', 'VeraIt.ttf'), ('Vera-BoldItalic', 'VeraBI.ttf')):
        pdfmetrics.registerFont(TTFont(name, str(base / file)))
    addMapping('Vera', 0, 0, 'Vera')
    addMapping('Vera', 1, 0, 'Vera-Bold')
    addMapping('Vera', 0, 1, 'Vera-Italic')
    addMapping('Vera', 1, 1, 'Vera-BoldItalic')
    _CMAP = set(pdfmetrics.getFont('Vera').face.charToGlyph.keys())
    _FONTS_READY = True


_CMAP: set = set()


def _drawable(value) -> str:
    """Every character the font can draw, or a stand-in for it."""
    out = []
    for ch in str(value or ''):
        ch = _FALLBACKS.get(ch, ch)
        if len(ch) > 1 or ch == '\n' or ord(ch) in _CMAP:
            out.append(ch)
        elif ord(ch) < 32:
            continue
        else:
            out.append('?' if ord(ch) < 0x1F000 else '')     # emoji simply go
    return ''.join(out)


def _pdf_text(value) -> str:
    """Text made safe for a reportlab Paragraph: drawable, XML escaped, and
    newlines as line breaks."""
    escaped = _drawable(value).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    return escaped.replace('\n', '<br/>')


def render_pdf(blocks: list[dict], meta: dict, images: dict) -> bytes:
    """`meta`: brand, campaign, name, pack, ref, op_date. `images`: id → bytes."""
    _register_fonts()
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image, KeepTogether, ListFlowable, ListItem, PageBreak,
                                    Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

    ink, muted, accent = (colors.Color(*(c / 255 for c in rgb)) for rgb in (_INK, _MUTED, _ACCENT))
    base = ParagraphStyle('base', fontName='Vera', fontSize=9.5, leading=13.5, textColor=ink)
    st = {
        'p': ParagraphStyle('p', parent=base, spaceAfter=6),
        'h1': ParagraphStyle('h1', parent=base, fontName='Vera-Bold', fontSize=15, leading=19,
                             spaceBefore=10, spaceAfter=4, keepWithNext=1),
        'h2': ParagraphStyle('h2', parent=base, fontName='Vera-Bold', fontSize=11, leading=14,
                             spaceBefore=8, spaceAfter=4, keepWithNext=1),
        'label': ParagraphStyle('label', parent=base, fontName='Vera-Bold', textColor=muted),
        'cell': ParagraphStyle('cell', parent=base, fontSize=9, leading=12),
        'head': ParagraphStyle('head', parent=base, fontName='Vera-Bold', fontSize=9, leading=12),
        'note': ParagraphStyle('note', parent=base, fontName='Vera-Italic', fontSize=10.5,
                               leading=15, leftIndent=8, borderPadding=(6, 8, 6, 8),
                               backColor=colors.Color(0.96, 0.95, 0.92), spaceBefore=4,
                               spaceAfter=10),
        'caption': ParagraphStyle('caption', parent=base, fontName='Vera-Italic', fontSize=8.5,
                                  leading=11, textColor=muted, spaceAfter=10),
        'cover_unit': ParagraphStyle('cu', parent=base, fontName='Vera-Bold', fontSize=13,
                                     alignment=TA_CENTER, textColor=muted, leading=17),
        'cover_title': ParagraphStyle('ct', parent=base, fontName='Vera-Bold', fontSize=28,
                                      leading=34, alignment=TA_CENTER, spaceBefore=18),
        'cover_sub': ParagraphStyle('cs', parent=base, fontSize=13, leading=18,
                                    alignment=TA_CENTER, textColor=ink, spaceBefore=6),
    }
    out = io.BytesIO()
    margin = 18 * mm
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=margin, rightMargin=margin,
                            topMargin=22 * mm, bottomMargin=18 * mm,
                            title=f"Operation {meta['name']} — {meta['pack']}",
                            author=meta.get('brand') or '')
    width = A4[0] - 2 * margin
    header = _drawable(f"{meta.get('brand') or ''}  ·  OPERATION {meta['name']}  ·  {meta['pack']}")

    def on_page(canvas, document):
        if document.page == 1:
            return
        canvas.saveState()
        canvas.setFont('Vera', 7.5)
        canvas.setFillColor(muted)
        canvas.drawString(margin, A4[1] - 13 * mm, header)
        canvas.drawRightString(A4[0] - margin, 10 * mm, f'{document.page}')
        canvas.setStrokeColor(accent)
        canvas.setLineWidth(0.6)
        canvas.line(margin, A4[1] - 15 * mm, A4[0] - margin, A4[1] - 15 * mm)
        canvas.restoreState()

    def picture(data: bytes, max_w: float, max_h: float):
        w, h = _size(data)
        scale = min(max_w / w, max_h / h)
        return Image(io.BytesIO(data), width=w * scale, height=h * scale)

    story: list = [Spacer(1, 40 * mm)]
    logo = logo_bytes()
    if logo:
        try:
            story += [picture(logo, 38 * mm, 38 * mm), Spacer(1, 8 * mm)]
        except Exception:
            pass
    if meta.get('brand'):
        story.append(Paragraph(_pdf_text(meta['brand']), st['cover_unit']))
    if meta.get('campaign'):
        story.append(Paragraph(_pdf_text(meta['campaign']), st['cover_sub']))
    story.append(Paragraph(_pdf_text(f"Operation {meta['name']}"), st['cover_title']))
    sub = meta['pack'] + (f"  ·  {meta['ref']}" if meta.get('ref') else '')
    story.append(Paragraph(_pdf_text(sub), st['cover_sub']))
    if meta.get('op_date'):
        story.append(Paragraph(_pdf_text(meta['op_date']), st['cover_sub']))
    rule = Table([['']], colWidths=[width * 0.5], rowHeights=[2])
    rule.setStyle(TableStyle([('LINEBELOW', (0, 0), (-1, -1), 1.2, accent)]))
    story += [Spacer(1, 6 * mm), rule, PageBreak()]

    grid = TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LINEBELOW', (0, 0), (-1, -1), 0.3, colors.Color(0.85, 0.87, 0.9)),
        ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 4), ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ])
    for block in blocks:
        kind = block['t']
        if kind == 'h1':
            story.append(Paragraph(_pdf_text(block['text']), st['h1']))
            line = Table([['']], colWidths=[width], rowHeights=[2])
            line.setStyle(TableStyle([('LINEABOVE', (0, 0), (-1, -1), 0.8, accent)]))
            line.keepWithNext = True       # a heading never ends a page alone
            story.append(line)
        elif kind == 'h2':
            story.append(Paragraph(_pdf_text(block['text']), st['h2']))
        elif kind == 'p':
            story.append(Paragraph(_pdf_text(block['text']), st['p']))
        elif kind == 'note':
            story.append(Paragraph(_pdf_text(block['text']), st['note']))
        elif kind == 'list':
            story.append(ListFlowable(
                [ListItem(Paragraph(_pdf_text(i), st['p']), leftIndent=12) for i in block['items']],
                bulletType='bullet', start='•', leftIndent=10, bulletFontName='Vera'))
        elif kind == 'kv':
            data = [[Paragraph(_pdf_text(k), st['label']), Paragraph(_pdf_text(v), st['cell'])]
                    for k, v in block['rows']]
            table = Table(data, colWidths=[width * 0.28, width * 0.72])
            table.setStyle(grid)
            story += [table, Spacer(1, 6)]
        elif kind == 'table':
            data = [[Paragraph(_pdf_text(h), st['head']) for h in block['head']]]
            data += [[Paragraph(_pdf_text(c), st['cell']) for c in row] for row in block['rows']]
            table = Table(data, colWidths=[width * f for f in block['widths']], repeatRows=1)
            style = TableStyle(grid.getCommands() + [
                ('BACKGROUND', (0, 0), (-1, 0), colors.Color(0.93, 0.94, 0.96)),
                ('LINEBELOW', (0, 0), (-1, 0), 0.8, accent)])
            table.setStyle(style)
            story += [table, Spacer(1, 8)]
        elif kind == 'img':
            data = images.get(block['id'])
            if not data:
                continue
            try:
                pic = picture(data, width if block.get('big') else width * 0.7,
                              200 * mm if block.get('big') else 110 * mm)
            except Exception:
                continue
            story.append(KeepTogether([pic, Paragraph(_pdf_text(block.get('caption')),
                                                      st['caption'])]))
        elif kind == 'break':
            story.append(PageBreak())
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------

def render_docx(blocks: list[dict], meta: dict, images: dict) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    ink, muted, accent = (RGBColor(*rgb) for rgb in (_INK, _MUTED, _ACCENT))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ('left_margin', 'right_margin'):
        setattr(section, side, Mm(18))
    section.top_margin, section.bottom_margin = Mm(20), Mm(18)
    width = section.page_width - section.left_margin - section.right_margin

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10)
    normal.font.color.rgb = ink
    for name, size in (('Heading 1', 15), ('Heading 2', 11.5)):
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ink

    def rule_under(paragraph, colour='B8973F', size=8):
        p_pr = paragraph._p.get_or_add_pPr()
        border = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        for k, v in (('w:val', 'single'), ('w:sz', str(size)), ('w:space', '4'),
                     ('w:color', colour)):
            bottom.set(qn(k), v)
        border.append(bottom)
        p_pr.append(border)

    def shade(cell, fill='EEF1F6'):
        tc_pr = cell._tc.get_or_add_tcPr()
        el = OxmlElement('w:shd')
        el.set(qn('w:val'), 'clear')
        el.set(qn('w:color'), 'auto')
        el.set(qn('w:fill'), fill)
        tc_pr.append(el)

    def write(cell_or_par, value, bold=False, italic=False, colour=None, size=None):
        paragraph = cell_or_par.paragraphs[0] if hasattr(cell_or_par, 'paragraphs') else cell_or_par
        for i, part in enumerate(str(value or '').split('\n')):
            run = paragraph.add_run(part)
            run.bold, run.italic = bold, italic
            if colour is not None:
                run.font.color.rgb = colour
            if size:
                run.font.size = Pt(size)
            if i < len(str(value or '').split('\n')) - 1:
                run.add_break()
        return paragraph

    def add_picture(data: bytes, max_w, max_h):
        data = _png_or_jpeg(data)
        w, h = _size(data)
        scale = min(max_w / w, max_h / h)
        doc.add_picture(io.BytesIO(data), width=int(w * scale))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.keep_with_next = True   # with its caption

    # Cover
    for _ in range(5):
        doc.add_paragraph()
    logo = logo_bytes()
    if logo:
        try:
            add_picture(logo, Mm(38), Mm(38))
        except Exception:
            pass
    for value, size, colour, bold in ((meta.get('brand'), 14, muted, True),
                                      (meta.get('campaign'), 13, ink, False),
                                      (f"Operation {meta['name']}", 30, ink, True),
                                      (meta['pack'] + (f"  ·  {meta['ref']}" if meta.get('ref') else ''),
                                       14, ink, False),
                                      (meta.get('op_date'), 12, muted, False)):
        if value:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            write(p, value, bold=bold, colour=colour, size=size)
    rule_under(doc.paragraphs[-1], size=12)

    # One section: the cover is its first page, which has its own (empty)
    # header, and every page after it carries the running header and number.
    section.different_first_page_header_footer = True
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    body = section
    head = body.header.paragraphs[0]
    write(head, f"{meta.get('brand') or ''}  ·  OPERATION {meta['name']}  ·  {meta['pack']}",
          colour=muted, size=8)
    rule_under(head, size=4)
    foot = body.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = foot.add_run()
    for tag, text_ in (('begin', None), (None, 'PAGE'), ('end', None)):
        if tag:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), tag)
        else:
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = text_
        run._r.append(el)

    def table(head_row, rows, widths, label_col=False):
        t = doc.add_table(rows=0, cols=len(widths))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = doc.styles['Table Grid']
        t.autofit = False
        for column, fraction in zip(t.columns, widths):
            column.width = int(width * fraction)
        if head_row:
            cells = t.add_row().cells
            for cell, value in zip(cells, head_row):
                write(cell, value, bold=True)
                shade(cell)
        for row in rows:
            cells = t.add_row().cells
            for i, (cell, value) in enumerate(zip(cells, row)):
                write(cell, value, bold=label_col and i == 0,
                      colour=muted if label_col and i == 0 else None)
        for row in t.rows:
            for cell, fraction in zip(row.cells, widths):
                cell.width = int(width * fraction)
        doc.add_paragraph()

    break_next = False
    for block in blocks:
        kind = block['t']
        if kind == 'break':
            # Carried onto the next heading as "page break before", not an
            # empty paragraph: a break paragraph landing at the foot of a full
            # page is what makes Word show a blank page.
            break_next = True
            continue
        if kind == 'h1':
            heading = doc.add_heading(block['text'], level=1)
            heading.paragraph_format.page_break_before = break_next
            heading.paragraph_format.keep_with_next = True
            break_next = False
            rule_under(heading)
        elif kind == 'h2':
            doc.add_heading(block['text'], level=2).paragraph_format.keep_with_next = True
        elif kind == 'p':
            write(doc.add_paragraph(), block['text'])
        elif kind == 'note':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Mm(5)
            write(p, block['text'], italic=True, size=11)
        elif kind == 'list':
            for item in block['items']:
                write(doc.add_paragraph(style='List Bullet'), item)
        elif kind == 'kv':
            table(None, block['rows'], [0.28, 0.72], label_col=True)
        elif kind == 'table':
            table(block['head'], block['rows'], block['widths'])
        elif kind == 'img':
            data = images.get(block['id'])
            if not data:
                continue
            try:
                add_picture(data, width if block.get('big') else int(width * 0.7),
                            Mm(200) if block.get('big') else Mm(110))
            except Exception:
                continue
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            write(p, block.get('caption'), italic=True, colour=muted, size=8.5)

    core = doc.core_properties
    core.title = f"Operation {meta['name']} — {meta['pack']}"
    core.author = meta.get('brand') or ''
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def op_date_label(value) -> str:
    try:
        return date.fromisoformat(str(value)).strftime('%A %d %B %Y')
    except (TypeError, ValueError):
        return ''


MEDIA = {'pdf': 'application/pdf',
         'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}
